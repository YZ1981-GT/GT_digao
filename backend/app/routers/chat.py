"""聊天路由层。

注册到 /api/chat 前缀，提供流式聊天等端点。
"""
import csv
import io
import json
import logging
import os
import shutil
import tempfile
from datetime import datetime

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from ..models.chat_schemas import (
    ChatStreamRequest, ChatUploadResponse, SpeechToTextResponse,
    NoteCreateRequest, NotesListResponse, NoteGroup, NoteItem,
    CleanupSuggestRequest, CleanupSuggestResponse, CleanupSuggestion,
    ExportWordRequest, PolishRequest,
)
from ..services.chat_service import chat_service
from ..services.knowledge_service import knowledge_service
from ..services.ocr_service import OCRService
from ..services.openai_service import OpenAIService
from ..services.word_service import WordExportService
from ..utils.docx_to_md import convert_docx_to_md
from ..utils.sse import sse_response, sse_with_heartbeat

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/chat", tags=["chat"])


@router.post("/stream")
async def chat_stream(request: ChatStreamRequest):
    """流式聊天端点。

    1. 调用 chat_service.build_messages 构建含 RAG 注入的 messages
    2. 调用 openai_service.stream_chat_completion 获取 SSE 流
    3. 收集完整回复文本，用于计算 meta 数据
    4. 流结束后发送 [DONE] 和 meta 数据行
    5. 使用 sse_with_heartbeat 包装防止连接断开
    """
    try:
        openai_service = OpenAIService()
        if not openai_service.api_key:
            raise HTTPException(status_code=400, detail="请先配置 AI API 密钥")

        # 将 Pydantic 模型转为 dict 列表
        messages_dicts = [m.model_dump() for m in request.messages]

        # 构建含 RAG 注入的完整 messages
        built_messages = await chat_service.build_messages(
            messages=messages_dicts,
            knowledge_library_ids=request.knowledge_library_ids,
            document_context=request.document_context,
            image_ocr_context=request.image_ocr_context,
        )

        async def generate():
            full_reply = ""
            try:
                async for chunk in openai_service.stream_chat_completion(
                    built_messages, temperature=0.7
                ):
                    full_reply += chunk
                    yield f"data: {chunk}\n\n"
            except Exception as e:
                logger.error("聊天流式生成失败: %s", e, exc_info=True)
                yield f"data: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"

            # 文本流结束标记
            yield "data: [DONE]\n\n"

            # 构建 meta 数据
            knowledge_refs = request.knowledge_library_ids or []
            suggest_save_note = chat_service.should_suggest_save_note(full_reply)
            meta = {
                "meta": {
                    "knowledge_refs": knowledge_refs,
                    "suggest_save_note": suggest_save_note,
                }
            }
            yield f"data: {json.dumps(meta, ensure_ascii=False)}\n\n"

        return sse_response(sse_with_heartbeat(generate()))

    except HTTPException:
        raise
    except Exception as e:
        logger.error("聊天请求处理失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"聊天请求处理失败: {str(e)}")


# ─── 文件大小上限 20MB ───
_MAX_UPLOAD_SIZE = 20 * 1024 * 1024

# 支持的文件扩展名
_SUPPORTED_EXTENSIONS = {
    '.pdf', '.docx', '.doc',
    '.xlsx', '.xls',
    '.csv',
    '.txt', '.md',
    '.png', '.jpg', '.jpeg', '.bmp', '.webp',
}

ocr_service = OCRService()


@router.post("/upload", response_model=ChatUploadResponse)
async def chat_upload(file: UploadFile = File(...)):
    """文档/图片上传端点。

    接收单个文件，根据扩展名提取文本内容：
    - PDF: OCR smart_parse
    - Word (.docx): docx_to_md 转换
    - Excel (.xlsx/.xls): openpyxl 读取
    - CSV: csv 模块读取
    - TXT/Markdown: 直接读取
    - 图片: OCR 识别
    """
    filename = file.filename or "unknown"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in _SUPPORTED_EXTENSIONS:
        return ChatUploadResponse(
            success=False,
            filename=filename,
            size=0,
            error=f"不支持的文件格式: {ext}，支持: {', '.join(sorted(_SUPPORTED_EXTENSIONS))}",
        )

    tmp_path = ""
    try:
        # 保存到临时文件
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            shutil.copyfileobj(file.file, tmp)
            tmp_path = tmp.name

        file_size = os.path.getsize(tmp_path)
        if file_size > _MAX_UPLOAD_SIZE:
            return ChatUploadResponse(
                success=False,
                filename=filename,
                size=file_size,
                error=f"文件大小 {file_size / 1024 / 1024:.1f}MB 超过限制（最大 20MB）",
            )

        content = ""
        ocr_method = ""

        # ─── 纯文本 ───
        if ext in ('.txt', '.md'):
            with open(tmp_path, 'r', encoding='utf-8-sig', errors='ignore') as f:
                content = f.read()

        # ─── CSV ───
        elif ext == '.csv':
            content = _parse_csv(tmp_path)

        # ─── 图片 → OCR ───
        elif OCRService.is_image_file(ext):
            content, ocr_method = await ocr_service.smart_parse(
                tmp_path, filename, ext
            )
            if not content.strip():
                return ChatUploadResponse(
                    success=False,
                    filename=filename,
                    size=file_size,
                    error="OCR 未能从图片中识别出文本内容，请确认图片清晰度",
                )

        # ─── PDF → OCR smart_parse ───
        elif ext == '.pdf':
            content, ocr_method = await ocr_service.smart_parse(
                tmp_path, filename, ext
            )
            if not content.strip():
                return ChatUploadResponse(
                    success=False,
                    filename=filename,
                    size=file_size,
                    error="未能从 PDF 中提取文本内容",
                )

        # ─── Word (.docx) → docx_to_md ───
        elif ext == '.docx':
            md_output = convert_docx_to_md(tmp_path)
            with open(md_output, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            # 清理生成的 md 文件
            if os.path.exists(md_output) and md_output != tmp_path:
                os.remove(md_output)

        # ─── Word (.doc) → 通过 OCR 服务处理 ───
        elif ext == '.doc':
            content, ocr_method = await ocr_service.smart_parse(
                tmp_path, filename, ext
            )

        # ─── Excel (.xlsx/.xls) → openpyxl ───
        elif ext in ('.xlsx', '.xls'):
            content = _parse_excel(tmp_path)

        return ChatUploadResponse(
            success=True,
            filename=filename,
            size=file_size,
            content=content,
            ocr_method=ocr_method,
        )

    except Exception as e:
        logger.error("文件上传处理失败: %s", e, exc_info=True)
        return ChatUploadResponse(
            success=False,
            filename=filename,
            size=0,
            error=f"文件处理失败: {str(e)}",
        )
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)


def _parse_csv(file_path: str) -> str:
    """读取 CSV 文件并转为文本表格。"""
    lines = []
    with open(file_path, 'r', encoding='utf-8-sig', errors='ignore') as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(' | '.join(row))
    return '\n'.join(lines)


def _parse_excel(file_path: str) -> str:
    """使用 openpyxl 读取 Excel 文件并转为文本。"""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_lines = [f"## {sheet_name}"]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else '' for c in row]
            if any(cells):  # 跳过全空行
                sheet_lines.append(' | '.join(cells))
        parts.append('\n'.join(sheet_lines))
    wb.close()
    return '\n\n'.join(parts)


# ─── 语音识别端点 ───

@router.post("/speech-to-text", response_model=SpeechToTextResponse)
async def chat_speech_to_text(audio: UploadFile = File(...)):
    """语音识别端点。

    接收 audio multipart 文件（WebM/WAV/MP3），调用 OpenAI 兼容的
    Whisper API（POST {base_url}/audio/transcriptions）进行语音识别。
    """
    import httpx
    from ..utils.config_manager import config_manager

    tmp_path = ""
    try:
        # 保存到临时文件
        ext = os.path.splitext(audio.filename or "recording.webm")[1] or ".webm"
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            shutil.copyfileobj(audio.file, tmp)
            tmp_path = tmp.name

        # 获取 API 配置
        cfg = config_manager.get_active_provider_config()
        api_key = cfg.get("api_key", "")
        base_url = (cfg.get("base_url", "") or "").rstrip("/")

        if not api_key:
            return SpeechToTextResponse(
                success=False, text="", error="请先配置 AI API 密钥"
            )
        if not base_url:
            return SpeechToTextResponse(
                success=False, text="", error="请先配置 API Base URL"
            )

        # 调用 Whisper API
        transcription_url = f"{base_url}/audio/transcriptions"

        async with httpx.AsyncClient(timeout=60.0) as client:
            with open(tmp_path, "rb") as f:
                resp = await client.post(
                    transcription_url,
                    headers={"Authorization": f"Bearer {api_key}"},
                    files={"file": (audio.filename or "recording.webm", f)},
                    data={"model": "whisper-1"},
                )

        if resp.status_code != 200:
            error_detail = resp.text[:200] if resp.text else f"HTTP {resp.status_code}"
            logger.error("Whisper API 调用失败: %s", error_detail)
            return SpeechToTextResponse(
                success=False, text="", error=f"语音识别失败: {error_detail}"
            )

        result = resp.json()
        text = result.get("text", "")
        return SpeechToTextResponse(success=True, text=text)

    except httpx.TimeoutException:
        logger.error("Whisper API 超时")
        return SpeechToTextResponse(
            success=False, text="", error="语音识别超时，请重试"
        )
    except Exception as e:
        logger.error("语音识别失败: %s", e, exc_info=True)
        return SpeechToTextResponse(
            success=False, text="", error=f"语音识别失败: {str(e)}"
        )
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)


# ─── 笔记 CRUD 端点 ───


@router.post("/notes")
async def create_note(request: NoteCreateRequest):
    """保存笔记到知识库。

    调用 knowledge_service.add_note 按日期子文件夹保存。
    """
    try:
        result = knowledge_service.add_note(
            title=request.title,
            content=request.content,
            date=request.date,
        )
        return {"success": True, **result}
    except Exception as e:
        logger.error("保存笔记失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"保存笔记失败: {str(e)}")


@router.get("/notes", response_model=NotesListResponse)
async def list_notes():
    """获取笔记列表，按日期分组返回。"""
    try:
        groups_raw = knowledge_service.get_notes_grouped()

        total_count = 0
        total_size = 0
        groups = []
        for g in groups_raw:
            notes = [
                NoteItem(
                    id=n['id'],
                    title=n['title'],
                    created_at=n['created_at'],
                    size=n['size'],
                )
                for n in g['notes']
            ]
            total_count += len(notes)
            total_size += sum(n.size for n in notes)
            groups.append(NoteGroup(date=g['date'], notes=notes))

        return NotesListResponse(
            groups=groups,
            total_count=total_count,
            total_size=total_size,
        )
    except Exception as e:
        logger.error("获取笔记列表失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"获取笔记列表失败: {str(e)}")


@router.delete("/notes/{doc_id}")
async def delete_note(doc_id: str):
    """删除笔记。"""
    try:
        success = knowledge_service.delete_note(doc_id)
        if not success:
            raise HTTPException(status_code=404, detail="笔记不存在")
        return {"success": True, "message": "笔记已删除"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error("删除笔记失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"删除笔记失败: {str(e)}")


@router.post("/notes/cleanup-suggest", response_model=CleanupSuggestResponse)
async def cleanup_suggest(request: CleanupSuggestRequest):
    """AI 笔记清理建议。

    1. 检查笔记数量/大小是否超过阈值
    2. 超过阈值时调用 LLM 分析笔记列表，生成清理建议
    3. 返回建议删除的笔记列表及原因
    """
    try:
        # 获取笔记列表
        groups_raw = knowledge_service.get_notes_grouped()

        total_count = 0
        total_size = 0
        all_notes: list[dict] = []
        for g in groups_raw:
            for n in g['notes']:
                total_count += 1
                total_size += n['size']
                all_notes.append({
                    'id': n['id'],
                    'title': n['title'],
                    'created_at': n['created_at'],
                    'size': n['size'],
                    'date': g['date'],
                })

        # 检查是否超过阈值
        threshold_size_bytes = int(request.threshold_size_mb * 1024 * 1024)
        if total_count <= request.threshold_count and total_size <= threshold_size_bytes:
            return CleanupSuggestResponse(suggestions=[])

        # 构建 LLM prompt，让 AI 分析哪些笔记可以清理
        notes_desc = "\n".join(
            f"- id: {n['id']}, 标题: {n['title']}, 创建日期: {n['created_at']}, "
            f"大小: {n['size']}字节"
            for n in all_notes
        )

        prompt = (
            f"以下是用户的笔记库，共 {total_count} 条笔记，"
            f"总大小 {total_size / 1024 / 1024:.1f}MB。\n"
            f"请分析这些笔记，识别可能重复、过时或低价值的笔记，"
            f"给出清理建议。\n\n"
            f"笔记列表：\n{notes_desc}\n\n"
            f"请返回 JSON 格式，结构如下：\n"
            f'{{"suggestions": [\n'
            f'  {{"doc_id": "笔记id", "title": "笔记标题", '
            f'"created_at": "创建时间", "reason": "建议删除原因"}}\n'
            f"]}}\n\n"
            f"只返回 JSON，不要其他文字。如果没有需要清理的笔记，返回空列表。"
        )

        messages = [
            {"role": "system", "content": "你是一个智能笔记管理助手，帮助用户整理和清理笔记库。"},
            {"role": "user", "content": prompt},
        ]

        # 调用 LLM（非流式，收集完整响应）
        openai_svc = OpenAIService()
        if not openai_svc.api_key:
            raise HTTPException(status_code=400, detail="请先配置 AI API 密钥")

        full_response = await openai_svc._collect_stream_text(messages, temperature=0.3)

        # 解析 LLM 返回的 JSON
        from ..utils.json_util import check_json
        suggestions: list[CleanupSuggestion] = []

        # 尝试从响应中提取 JSON
        response_text = full_response.strip()
        # 去除可能的 markdown 代码块包裹
        if response_text.startswith("```"):
            lines = response_text.split("\n")
            # 去掉首尾的 ``` 行
            lines = [l for l in lines if not l.strip().startswith("```")]
            response_text = "\n".join(lines).strip()

        try:
            parsed = json.loads(response_text)
            raw_suggestions = parsed.get("suggestions", [])
            # 验证每条建议的 doc_id 确实存在
            valid_ids = {n['id'] for n in all_notes}
            for s in raw_suggestions:
                doc_id = s.get("doc_id", "")
                if doc_id in valid_ids:
                    suggestions.append(CleanupSuggestion(
                        doc_id=doc_id,
                        title=s.get("title", ""),
                        created_at=s.get("created_at", ""),
                        reason=s.get("reason", ""),
                        last_rag_reference=s.get("last_rag_reference"),
                    ))
        except (json.JSONDecodeError, AttributeError):
            logger.warning("LLM 清理建议响应解析失败: %s", response_text[:200])

        return CleanupSuggestResponse(suggestions=suggestions)

    except HTTPException:
        raise
    except Exception as e:
        logger.error("生成清理建议失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"生成清理建议失败: {str(e)}")


# ─── Word 导出端点 ───


@router.post("/export-word")
async def export_word(request: ExportWordRequest):
    """将 Markdown 内容导出为 Word 文件流。"""
    try:
        svc = WordExportService()
        svc._add_markdown_content(request.content)

        buf = io.BytesIO()
        svc.doc.save(buf)
        buf.seek(0)

        # 文件名
        if request.filename:
            fname = request.filename
        else:
            now = datetime.now()
            fname = f"聊天记录_{now.strftime('%Y%m%d_%H%M%S')}"

        from urllib.parse import quote
        encoded_fname = quote(f"{fname}.docx")

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_fname}",
            },
        )
    except Exception as e:
        logger.error("Word 导出失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Word 导出失败: {str(e)}")


# ─── AI 润色端点 ───


@router.post("/polish")
async def polish_content(request: PolishRequest):
    """AI 润色：SSE 流式返回润色后的内容。"""
    try:
        openai_service = OpenAIService()
        if not openai_service.api_key:
            raise HTTPException(status_code=400, detail="请先配置 AI API 密钥")

        messages = [
            {
                "role": "system",
                "content": (
                    "你是一个专业的文字润色助手。请对用户提供的内容进行润色：\n"
                    "1. 去除对话格式（如'提问：''回复：'），统一为连贯的文章\n"
                    "2. 统一文风，使行文流畅自然\n"
                    "3. 添加适当的过渡语句和段落衔接\n"
                    "4. 保留所有关键信息和专业术语\n"
                    "5. 输出 Markdown 格式"
                ),
            },
            {"role": "user", "content": f"请润色以下内容：\n\n{request.content}"},
        ]

        async def generate():
            try:
                async for chunk in openai_service.stream_chat_completion(
                    messages, temperature=0.5
                ):
                    yield f"data: {chunk}\n\n"
            except Exception as e:
                logger.error("AI 润色流式生成失败: %s", e, exc_info=True)
                yield f"data: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"
            yield "data: [DONE]\n\n"

        return sse_response(sse_with_heartbeat(generate()))

    except HTTPException:
        raise
    except Exception as e:
        logger.error("AI 润色请求处理失败: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=f"AI 润色请求处理失败: {str(e)}")
